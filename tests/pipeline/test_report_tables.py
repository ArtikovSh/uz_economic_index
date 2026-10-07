import ast
from contextlib import closing
from copy import deepcopy
from pathlib import Path
import re
from types import SimpleNamespace

import gspread
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string
import pandas as pd
import pytest

import excel_exporter
import post_text
import report_tables
import sheets_sync
import store


ROOT = Path(__file__).resolve().parents[2]
DAILY = ["Sana", "Hafta kuni", "Jami xabarlar", "Reklama emas", "Iqtisodiy", "Ijobiy", "Neytral",
         "Salbiy", "EAI, %", "ESI", "Kanallar", "Izoh"]
PERIOD = ["Davr", "Boshlanish", "Tugash", "Kunlar", "Kunlar (jami)", "Jami xabarlar", "Reklama emas",
          "Iqtisodiy", "Ijobiy", "Neytral", "Salbiy", "EAI, %", "ESI", "EAI farqi", "ESI farqi", "Izoh"]
TOPICS = ["Sana", "Mavzu", "Xabarlar", "Ijobiy", "Neytral", "Salbiy", "Mavzu ESI", "ESI'ga hissa"]
POSTS = ["Sana", "Vaqt", "Kanal", "Sarlavha", "Matn boshi", "Mavzu", "Ohang", "Indeksda", "Sabab",
         "Ko'rishlar", "Forwardlar", "Havola"]
HEADERS = {"Kunlik": DAILY, "Haftalik": PERIOD, "Oylik": PERIOD, "Choraklik": PERIOD,
           "Yillik": PERIOD, "Mavzular": TOPICS, "Xabarlar": POSTS}


@pytest.fixture(scope="module")
def real_data():
    return (pd.read_csv(ROOT / "data/indices.csv").reindex(columns=store.INDEX_COLS),
            pd.read_csv(ROOT / "data/messages.csv").reindex(columns=store.LEDGER_COLS))


@pytest.fixture(scope="module")
def real_tables(real_data):
    return report_tables.build_tables(*real_data)


def test_real_tables_schema_counts_and_types(real_data, real_tables):
    indices, ledger = real_data
    assert list(real_tables) == list(HEADERS)
    for title, frame in real_tables.items():
        assert list(frame.columns) == HEADERS[title]
        for column in frame:
            for value in frame[column]:
                assert value is None or not pd.isna(value)
                if value is not None and column in {"Jami xabarlar", "Reklama emas", "Iqtisodiy", "Ijobiy",
                                                   "Neytral", "Salbiy", "Kanallar", "Xabarlar", "Kunlar",
                                                   "Kunlar (jami)", "Ko'rishlar", "Forwardlar"}:
                    assert type(value) is int
                if value is not None and column in {"EAI, %", "ESI", "Mavzu ESI", "EAI farqi", "ESI farqi"}:
                    assert type(value) is float and value == round(value, 1)
    daily = real_tables["Kunlik"]
    assert len(daily) == (indices["period_type"] == "kun").sum()
    assert daily["Sana"].is_monotonic_increasing
    channels = ledger.groupby(ledger["date_local"].str[:10])["channel"].nunique()
    assert daily["Kanallar"].tolist() == [channels[d] for d in daily["Sana"]]
    assert daily["Hafta kuni"].iloc[0] == "chorshanba"
    assert len(real_tables["Xabarlar"]) == len(ledger)


def test_real_topic_contributions_and_sort(real_tables):
    topics = real_tables["Mavzular"]
    daily = real_tables["Kunlik"].set_index("Sana")
    for day, group in topics.groupby("Sana"):
        assert sum(group["ESI'ga hissa"]) == pytest.approx(daily.loc[day, "ESI"], abs=0.050000001)
        assert group["Xabarlar"].is_monotonic_decreasing
        assert sum(group["Xabarlar"]) == daily.loc[day, "Iqtisodiy"]
        for row in group.to_dict("records"):
            balance = row["Ijobiy"] - row["Salbiy"]
            assert row["Mavzu ESI"] == round(100 * balance / row["Xabarlar"], 1)
            assert row["ESI'ga hissa"] == pytest.approx(100 * balance / daily.loc[day, "Iqtisodiy"])


def test_real_post_labels_reasons_and_order(real_data, real_tables):
    _, ledger = real_data
    posts = real_tables["Xabarlar"]
    assert list(zip(posts["Sana"], posts["Vaqt"])) == sorted(zip(posts["Sana"], posts["Vaqt"]))
    original = ledger.set_index(["channel", "message_id"])
    for row in posts.to_dict("records"):
        post = original.loc[row["Kanal"], int(row["Havola"].rsplit("/", 1)[1])]
        assert row["Havola"].startswith("https://t.me/" + row["Kanal"].lstrip("@") + "/")
        assert row["Indeksda"] == ("ha" if post["econ"] else "yo'q")
        if row["Indeksda"] == "ha":
            assert row["Sabab"] is None
            assert row["Ohang"] == {1: "ijobiy", 0: "neytral", -1: "salbiy"}[post["tone"]]
        else:
            reason = ("reklama" if post["is_ad"] else "dayjest" if post["is_digest"] else
                      "xorijiy" if post["is_foreign"] else "iqtisodiy emas")
            assert row["Sabab"] == reason
            assert row["Ohang"] is None
        if post["primary_topic"] == "non_economic":
            assert row["Mavzu"] == "Iqtisodiy emas"


def _index(kind, period, start, end, eai=40.0, esi=20.0):
    return dict(period_type=kind, period=period, start=start, end=end, days=1, days_expected=1,
                posts=10, nonad=10, econ=4, pos=2, neu=1, neg=1, EAI=eai, ESI=esi, note=None)


@pytest.mark.parametrize("kind,period,start,end,title,label", [
    ("hafta", "2026-W40", "2026-09-28", "2026-10-04", "Haftalik", "40-hafta (28.09–04.10.2026)"),
    ("oy", "2026-09", "2026-09-01", "2026-09-30", "Oylik", "Sentabr 2026"),
    ("chorak", "2026-Q3", "2026-07-01", "2026-09-30", "Choraklik", "2026, III chorak"),
    ("yil", "2026", "2026-01-01", "2026-12-31", "Yillik", "2026"),
    ("hafta", "2026-W53", "2026-12-28", "2027-01-03", "Haftalik", "53-hafta (28.12.2026–03.01.2027)"),
])
def test_period_labels_and_missing_previous(kind, period, start, end, title, label):
    indices = pd.DataFrame([_index(kind, period, start, end)])
    result = report_tables.build_tables(indices, pd.DataFrame(columns=store.LEDGER_COLS))[title].iloc[0]
    assert result["Davr"] == label
    assert result["EAI farqi"] is None and result["ESI farqi"] is None


@pytest.mark.parametrize("kind,title,starts", [
    ("hafta", "Haftalik", ["2025-12-22", "2025-12-29", "2026-01-12"]),
    ("oy", "Oylik", ["2025-12-01", "2026-01-01", "2026-03-01"]),
    ("chorak", "Choraklik", ["2025-10-01", "2026-01-01", "2026-07-01"]),
    ("yil", "Yillik", ["2024-01-01", "2025-01-01", "2027-01-01"]),
])
def test_deltas_use_immediate_calendar_period(kind, title, starts):
    rows = [_index(kind, str(i), start, start, 40.0 + i * 1.3, 20.0 - i * 2.7)
            for i, start in enumerate(starts)]
    # Another period type at the missing start must not be used as the predecessor.
    rows.append(_index("kun", "2026-01-01", "2026-01-01", "2026-01-01"))
    frame = report_tables.build_tables(pd.DataFrame(rows[::-1]), pd.DataFrame(columns=store.LEDGER_COLS))[title]
    assert frame["EAI farqi"].tolist() == [None, 1.3, None]
    assert frame["ESI farqi"].tolist() == [None, -2.7, None]


def test_empty_tables_and_inputs_are_unchanged(real_data):
    empty = report_tables.build_tables(pd.DataFrame(columns=store.INDEX_COLS), pd.DataFrame(columns=store.LEDGER_COLS))
    assert all(frame.empty and list(frame.columns) == HEADERS[name] for name, frame in empty.items())
    indices, ledger = real_data
    before = indices.copy(deep=True), ledger.copy(deep=True)
    report_tables.build_tables(indices, ledger)
    pd.testing.assert_frame_equal(indices, before[0])
    pd.testing.assert_frame_equal(ledger, before[1])


def test_api_post_parts_identical_for_300_real_posts(real_data):
    # Execute the API's actual pure helpers without importing its DB/bot environment.
    source = (ROOT / "app/api/index.py").read_text(encoding="utf-8")
    names = {"_MD_LINK", "_URL", "_EMOJI", "_BOILERPLATE", "_clean_lines", "_cut", "post_parts"}
    nodes = [node for node in ast.parse(source).body if
             (isinstance(node, ast.FunctionDef) and node.name in names) or
             (isinstance(node, ast.Assign) and any(isinstance(t, ast.Name) and t.id in names for t in node.targets))]
    namespace = {"re": re}
    exec(compile(ast.Module(body=nodes, type_ignores=[]), "app/api/index.py", "exec"), namespace)
    for node in nodes:
        assert ast.get_source_segment(source, node) in (ROOT / "post_text.py").read_text(encoding="utf-8")
    for raw in real_data[1]["raw_text"].sample(n=300, random_state=5):
        assert post_text.post_parts(raw, 200) == namespace["post_parts"](raw, 200)


def test_excel_layout_and_literal_text(tmp_path, monkeypatch, real_data):
    indices, ledger = real_data
    ledger = ledger.head(12).iloc[::-1].copy()
    ledger.loc[ledger.index[0], "raw_text"] = "=1+1\nMatn"
    monkeypatch.setattr(excel_exporter, "OUTPUT_DIR", str(tmp_path))
    path = excel_exporter.export_results(indices, ledger, pending_count=3)
    expected = report_tables.build_tables(indices, ledger)
    with closing(load_workbook(path)) as book:
        assert book.sheetnames == list(HEADERS)
        assert "Metodika" not in book.sheetnames
        for title, header in HEADERS.items():
            ws = book[title]
            cols = header + (["To'liq matn", "Model"] if title == "Xabarlar" else [])
            assert [cell.value for cell in ws[1]] == cols
            assert ws.freeze_panes == "A2"
            assert ws.auto_filter.ref == ws.dimensions
            assert all(cell.font.bold for cell in ws[1])
            assert ws.max_row == len(expected[title]) + 1
            for column, name in enumerate(cols, 1):
                if len(expected[title]) and ("EAI" in name or "ESI" in name):
                    assert ws.cell(2, column).number_format == "0.0"
            if len(expected[title]) and title != "Xabarlar":
                rules = [rule for rule_list in ws.conditional_formatting._cf_rules.values() for rule in rule_list]
                assert {rule.dxf.font.color.rgb[-6:] for rule in rules} == {"0E7490", "C2410C"}
        ws = book["Xabarlar"]
        sorted_posts = ledger.sort_values(["date_local", "channel", "message_id"])
        for i, post in enumerate(sorted_posts.to_dict("records"), 2):
            assert ws.cell(i, 13).value == post["raw_text"].replace("\r\n", "\n").replace("\r", "\n")
            assert ws.cell(i, 14).value == post["label_model"]
            assert ws.cell(i, 13).data_type == "s"
            assert ws.cell(i, 4).alignment.wrap_text
            assert ws.cell(i, 5).alignment.wrap_text
        assert ws.column_dimensions["D"].width >= 50


class FakeWorksheet:
    def __init__(self, sheet, title, rows, cols):
        self.sheet, self.title = sheet, title
        self.id = len(sheet.tabs) + 1
        self.row_count, self.col_count = rows, cols
        self.values, self.rules, self.appends = [], [], []
        self.clears = 0

    def row_values(self, row):
        self.sheet.calls += 1
        return deepcopy(self.values[row - 1] if len(self.values) >= row else [])

    def get(self, key_range):
        self.sheet.calls += 1
        match = re.fullmatch(r"([A-Z]+)2:([A-Z]+)", key_range)
        assert match
        first, last = (column_index_from_string(c) for c in match.groups())
        return [row[first - 1:last] for row in self.values[1:]]

    def clear(self):
        self.sheet.calls += 1
        self.clears += 1
        self.values = []

    def update(self, values, range_name, value_input_option):
        self.sheet.calls += 1
        assert range_name == "A1" and value_input_option == "RAW"
        self.values = deepcopy(values)

    def append_rows(self, values, value_input_option):
        self.sheet.calls += 1
        assert value_input_option == "RAW" and len(values) <= 2000
        assert len(self.values) + len(values) <= self.row_count
        self.appends.append(len(values))
        self.values.extend(deepcopy(values))


class FakeSpreadsheet:
    title = "Test report"

    def __init__(self):
        self.tabs, self.calls, self.batches = {}, 0, []

    def worksheets(self):
        self.calls += 1
        return list(self.tabs.values())

    def fetch_sheet_metadata(self, params):
        self.calls += 1
        return {"sheets": [{"properties": {"sheetId": ws.id}, "conditionalFormats": deepcopy(ws.rules)}
                           for ws in self.tabs.values()]}

    def add_worksheet(self, title, rows, cols):
        self.calls += 1
        ws = FakeWorksheet(self, title, rows, cols)
        self.tabs[title] = ws
        return ws

    def batch_update(self, body):
        self.calls += 1
        self.batches.append(deepcopy(body))
        for request in body["requests"]:
            kind, data = next(iter(request.items()))
            if kind == "updateSheetProperties":
                props = data["properties"]
                ws = next(ws for ws in self.tabs.values() if ws.id == props["sheetId"])
                ws.row_count = props["gridProperties"]["rowCount"]
                ws.col_count = props["gridProperties"]["columnCount"]
                assert props["gridProperties"]["frozenRowCount"] == 1
            elif kind == "deleteSheet":
                title = next(ws.title for ws in self.tabs.values() if ws.id == data["sheetId"])
                del self.tabs[title]
            elif kind in {"addConditionalFormatRule", "deleteConditionalFormatRule"}:
                sheet_id = data["rule"]["ranges"][0]["sheetId"] if "rule" in data else data["sheetId"]
                ws = next(ws for ws in self.tabs.values() if ws.id == sheet_id)
                if kind == "addConditionalFormatRule":
                    ws.rules.insert(data["index"], data["rule"])
                else:
                    ws.rules.pop(data["index"])


def _mock_google(monkeypatch, sheet, real_data):
    indices, ledger = real_data
    monkeypatch.setattr(sheets_sync, "SA_JSON", '{"client_email":"test@example.invalid"}')
    monkeypatch.setattr(sheets_sync, "SHEET_ID", "test-sheet")
    monkeypatch.setattr(gspread, "service_account_from_dict", lambda *a, **kw:
                        SimpleNamespace(open_by_key=lambda key: sheet))
    monkeypatch.setattr(sheets_sync, "load_indices", lambda: indices)
    monkeypatch.setattr(sheets_sync, "load_ledger", lambda: ledger)
    monkeypatch.setattr(sheets_sync, "load_pending", lambda: pd.DataFrame([{}, {}, {}]))


@pytest.mark.parametrize("full", [False, True])
def test_sheets_migration_and_chunking(monkeypatch, real_data, real_tables, full):
    sheet = FakeSpreadsheet()
    for name in sheets_sync.LEGACY_TABS:
        sheet.add_worksheet(name, 100, 20)
    old_posts = sheet.add_worksheet("Xabarlar", 6000, 18)
    old_posts.values = [["obsolete header"], ["stale"]]
    _mock_google(monkeypatch, sheet, real_data)
    sheet.calls = 0
    assert sheets_sync.main(full=full) == 0
    assert sheet.calls + 2 <= 60  # Include authentication/open_by_key overhead.
    assert list(sheet.tabs) == ["Xabarlar", *[n for n in HEADERS if n != "Xabarlar"], "Info"]
    assert not set(sheets_sync.LEGACY_TABS) & set(sheet.tabs)
    for name, frame in real_tables.items():
        assert sheet.tabs[name].values == [HEADERS[name], *sheets_sync._rows(frame)]
    assert old_posts.clears == 1
    assert old_posts.appends == [2000, 2000, len(real_data[1]) - 4000]
    info = dict(sheet.tabs["Info"].values[1:])
    assert info["Label versiyasi"] == sheets_sync.LLM_LABEL_VERSION
    assert info["Kun yakunlanishini kutayotgan xabarlar"] == 3
    assert len(sheet.batches) == 1
    requests = sheet.batches[0]["requests"]
    assert sum("setBasicFilter" in r for r in requests) == 8
    assert any(r.get("repeatCell", {}).get("cell", {}).get("userEnteredFormat", {}).get("numberFormat")
               == {"type": "NUMBER", "pattern": "0.0"} for r in requests)


def test_sheets_append_fills_holes_preserves_existing_and_is_idempotent(monkeypatch, real_data, real_tables):
    sheet = FakeSpreadsheet()
    _mock_google(monkeypatch, sheet, real_data)
    before = {}
    for name, frame in real_tables.items():
        ws = sheet.add_worksheet(name, 6000, len(frame.columns))
        rows = sheets_sync._rows(frame)
        ws.values = [list(frame.columns), *rows[-2:]]
        # Existing values are intentionally stale: normal sync must preserve them.
        if rows:
            ws.values[-1][-1 if name != "Xabarlar" else 3] = "existing value"
        before[name] = deepcopy(ws.values)
    sheet.calls = 0
    assert sheets_sync.main() == 0
    assert sheet.calls + 2 <= 60
    for name, frame in real_tables.items():
        ws = sheet.tabs[name]
        assert ws.clears == 0
        assert ws.values[:len(before[name])] == before[name]
        assert len(ws.values) == len(frame) + 1
    snapshot = {name: deepcopy(ws.values) for name, ws in sheet.tabs.items() if name != "Info"}
    rule_counts = {name: len(ws.rules) for name, ws in sheet.tabs.items()}
    assert sheets_sync.main() == 0
    assert snapshot == {name: ws.values for name, ws in sheet.tabs.items() if name != "Info"}
    assert rule_counts == {name: len(ws.rules) for name, ws in sheet.tabs.items()}
    assert sheets_sync.main(full=True) == 0
    for name, frame in real_tables.items():
        assert sheet.tabs[name].values == [list(frame.columns), *sheets_sync._rows(frame)]
        assert sheet.tabs[name].clears == 1


def test_sheets_empty_reports(monkeypatch):
    data = pd.DataFrame(columns=store.INDEX_COLS), pd.DataFrame(columns=store.LEDGER_COLS)
    sheet = FakeSpreadsheet()
    _mock_google(monkeypatch, sheet, data)
    assert sheets_sync.main(full=True) == 0
    for name in HEADERS:
        assert sheet.tabs[name].values == [HEADERS[name]]
